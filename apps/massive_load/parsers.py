"""Lectura de los archivos de carga masiva.

Se soportan CSV y XLSX, que son los dos formatos que exige la documentacion.

El parseo produce diccionarios con las cabeceras normalizadas; NO valida
contenido de negocio: de eso se encarga el servicio, fila a fila, para que un
dato incorrecto no tumbe la importacion entera.

DEC-33: ni el DER ni la documentacion funcional definen las columnas del
archivo. El contrato canonico se fija aqui, alineado con lo que
`IMPORT_BATCH_ROW` guarda como snapshot, y los archivos reales se adaptan a el
con `MASSIVE_LOAD_COLUMN_ALIASES` -- sin tocar codigo.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from django.conf import settings
from django.core.exceptions import ValidationError

# Columnas obligatorias del archivo.
COLUMNAS_OBLIGATORIAS: frozenset[str] = frozenset(
    {"identification_number", "first_name", "last_name", "pathology"}
)

# Columnas opcionales reconocidas. Cualquier otra se ignora en silencio: un
# archivo exportado de otro sistema suele traer columnas de mas.
COLUMNAS_OPCIONALES: frozenset[str] = frozenset(
    {
        "date_of_birth",
        "phone",
        "email",
        "address",
        "diagnosis_date",
        "is_primary",
        "notes",
    }
)

COLUMNAS_RECONOCIDAS = COLUMNAS_OBLIGATORIAS | COLUMNAS_OPCIONALES

# Tope defensivo: un archivo enorme no debe agotar la memoria del worker.
MAX_FILAS = 10_000


class ImportParseError(ValidationError):
    """El archivo no se puede leer o no tiene la estructura esperada.

    Es distinto de un error de fila: aqui no hay nada que importar, asi que se
    rechaza la subida entera con 400.
    """


def _plegar(nombre: Any) -> str:
    """Forma comparable de una cabecera: sin espacios sobrantes, en minusculas."""
    return str(nombre or "").strip().lower().replace(" ", "_")


def alias_configurados() -> dict[str, str]:
    """Mapeo `cabecera del archivo -> columna canonica` (DEC-33).

    Las claves se pliegan igual que las cabeceras del archivo, para que
    configurar `"Numero Identificacion"` funcione igual que
    `"numero_identificacion"`.
    """
    return {
        _plegar(origen): destino
        for origen, destino in getattr(
            settings, "MASSIVE_LOAD_COLUMN_ALIASES", {}
        ).items()
    }


def _normalizar_cabecera(nombre: Any) -> str:
    plegada = _plegar(nombre)
    return alias_configurados().get(plegada, plegada)


def _comprobar_cabeceras(cabeceras: list[str]) -> None:
    faltan = COLUMNAS_OBLIGATORIAS - set(cabeceras)
    if not faltan:
        return
    # Decir solo lo que falta obliga a adivinar como se llamaba la columna en el
    # archivo. Diciendo tambien lo que se encontro, la diferencia se ve de un
    # vistazo y se corrige con un alias (DEC-33).
    encontradas = ", ".join(sorted(c for c in cabeceras if c)) or "ninguna"
    raise ImportParseError(
        "Al archivo le faltan columnas obligatorias: "
        + ", ".join(sorted(faltan))
        + f". Columnas encontradas: {encontradas}. Si el archivo las nombra de "
        "otra forma, mapealas con MASSIVE_LOAD_COLUMN_ALIASES."
    )


def _limpiar(valor: Any) -> str | None:
    """Normaliza una celda a texto sin espacios, o `None` si esta vacia.

    Las fechas se convierten a ISO, NO con `str()`: openpyxl devuelve objetos
    `datetime` para las celdas de fecha, y `str()` las dejaria como
    "1980-05-14 00:00:00", que ninguna mascara de fecha admite. ISO ademas es
    serializable a JSON, que es como viajan las filas hasta la tarea Celery.
    """
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    return texto or None


def _fila(datos: dict[str, Any]) -> dict[str, Any]:
    """Se queda solo con las columnas reconocidas."""
    return {
        clave: _limpiar(valor)
        for clave, valor in datos.items()
        if clave in COLUMNAS_RECONOCIDAS
    }


def parse_csv(contenido: bytes) -> list[dict[str, Any]]:
    try:
        # `utf-8-sig` descarta el BOM que anade Excel al exportar a CSV.
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportParseError("El archivo CSV no esta codificado en UTF-8.") from exc

    lector = csv.DictReader(io.StringIO(texto))
    if not lector.fieldnames:
        raise ImportParseError("El archivo CSV esta vacio o no tiene cabecera.")

    lector.fieldnames = [_normalizar_cabecera(c) for c in lector.fieldnames]
    _comprobar_cabeceras(lector.fieldnames)

    filas = []
    for numero, datos in enumerate(lector, start=1):
        if numero > MAX_FILAS:
            raise ImportParseError(f"El archivo supera el maximo de {MAX_FILAS} filas.")
        filas.append(_fila(datos))
    return filas


def parse_xlsx(contenido: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        # `read_only` no carga la hoja entera en memoria.
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportParseError("El archivo XLSX no se puede leer.") from exc

    try:
        hoja = libro.active
        iterador = hoja.iter_rows(values_only=True)
        try:
            cabecera = [_normalizar_cabecera(c) for c in next(iterador)]
        except StopIteration as exc:
            raise ImportParseError("El archivo XLSX esta vacio.") from exc

        _comprobar_cabeceras(cabecera)

        filas = []
        for numero, valores in enumerate(iterador, start=1):
            if numero > MAX_FILAS:
                raise ImportParseError(
                    f"El archivo supera el maximo de {MAX_FILAS} filas."
                )
            if all(valor is None for valor in valores):
                # Excel suele dejar filas vacias al final de la hoja.
                continue
            filas.append(_fila(dict(zip(cabecera, valores, strict=False))))
        return filas
    finally:
        libro.close()


def parse(contenido: bytes, extension: str) -> list[dict[str, Any]]:
    """Lee el archivo segun su extension, ya validada por `core.validators`."""
    if extension == ".csv":
        return parse_csv(contenido)
    return parse_xlsx(contenido)
